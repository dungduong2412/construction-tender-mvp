import io
import os
import subprocess
import sys
import time
from pathlib import Path

import fitz
import httpx
import pytest
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parent.parent
BACKEND = ROOT / "backend"


def sample_pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Construction Tender UI test source")
    payload = document.tobytes()
    document.close()
    return payload


def start_app(database_path: Path, document_storage: Path):
    env = os.environ.copy()
    env["MOCK_PARSER"] = "true"
    env["MOCK_AI"] = "true"
    env["DATABASE_URL"] = f"sqlite+aiosqlite:///{database_path}"
    env["DOCUMENT_STORAGE_DIR"] = str(document_storage)
    proc = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "main:app", "--host", "127.0.0.1", "--port", "8000"],
        cwd=str(BACKEND), env=env, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
    )
    deadline = time.time() + 25
    while time.time() < deadline:
        try:
            if httpx.get("http://127.0.0.1:8000/api/parser/contract", timeout=1.5).status_code == 200:
                return proc
        except Exception:
            time.sleep(0.25)
    stdout, stderr = proc.communicate(timeout=5)
    raise RuntimeError(f"Backend failed to start. stdout={stdout} stderr={stderr}")


@pytest.fixture
def app_server(tmp_path):
    proc = start_app(tmp_path / "frontend-test.db", tmp_path / "documents")
    yield proc
    proc.terminate()
    try:
        proc.wait(timeout=10)
    except subprocess.TimeoutExpired:
        proc.kill()
        proc.wait(timeout=10)


def test_simplified_vietnamese_estimator_flow(app_server):
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1440, "height": 1200})
        try:
            page.goto("http://127.0.0.1:8000", wait_until="domcontentloaded")
            assert page.locator(".step[data-route]").count() == 3
            assert "Tải bảng tiên lượng" in page.locator('[data-route="upload"]').text_content()
            assert "Kiểm tra & Đối chiếu" in page.locator('[data-route="review"]').text_content()
            assert "Lập dự thầu" in page.locator('[data-route="estimate"]').text_content()
            assert page.locator('[data-route="formula"]').count() == 0
            assert page.locator('[data-route="calculation"]').count() == 0
            assert page.locator('[data-route="master"]').count() == 0

            page.set_input_files("#pdf-input", [{
                "name": "sample.pdf", "mimeType": "application/pdf", "buffer": sample_pdf_bytes(),
            }])
            page.locator("#upload-btn").click()
            page.wait_for_function(
                "() => document.getElementById('page-review').classList.contains('active')",
                timeout=45000,
            )

            assert page.locator("#review-tbody tr").count() > 0
            assert page.locator("#summary-bar .kpi").count() == 4
            page.locator("#pdf-image").wait_for(state="visible", timeout=10000)
            assert page.locator("#pdf-image").is_visible()
            assert "KS-002" in page.locator(
                "#review-tbody tr", has_text="Đo vẽ bình đồ tỷ lệ 1/500 vùng đồi núi"
            ).first.text_content()
            row = page.locator(
                "#review-tbody tr", has_text="Đo vẽ bình đồ tỷ lệ 1/500 vùng đồi núi"
            ).first
            row.click()
            assert "trang" in page.locator("#source-locator").text_content()

            row.locator(".btn-edit").click()
            page.locator("#modal-master-sel").select_option(index=1)
            page.locator("#modal-price-input").fill("250000")
            page.locator("#modal-price-source").fill("Biên bản duyệt đơn giá UAT")
            page.locator("#modal-price-effective").fill("2026-09-28")
            page.locator("#modal-price-approver").fill("uat-reviewer")
            page.locator("#modal-save").click()
            page.wait_for_timeout(1000)
            assert not page.locator("#modal-error").is_visible(), page.locator("#modal-error").text_content()
            page.wait_for_function(
                "() => !document.getElementById('modal-overlay').classList.contains('open')",
                timeout=20000,
            )

            page.locator('[data-route="estimate"]').click()
            page.wait_for_function("() => typeof currentSnapshot !== 'undefined' && currentSnapshot !== null", timeout=20000)
            assert page.locator("[data-estimate-tab]").count() == 2
            assert page.locator("#estimate-rows tr").count() > 0
            assert page.locator("#draft-warn").is_visible()
            assert "Chưa có đơn giá" in page.locator("#estimate-rows").text_content()
            assert page.locator("#add-row").is_visible()

            job_id = page.evaluate("currentJobId")
            snapshot_id = page.evaluate("currentSnapshot.snapshot_id")
            for view, expected_sheet in [("approval", "Phê duyệt nội bộ"), ("tender", "Dự thầu")]:
                response = page.request.get(
                    f"http://127.0.0.1:8000/api/export/{job_id}/estimate.xlsx"
                    f"?view={view}&snapshot_id={snapshot_id}"
                )
                assert response.status == 200
                assert response.headers["x-export-status"] == "draft-incomplete"
                workbook = load_workbook(io.BytesIO(response.body()), data_only=False)
                assert workbook.sheetnames == [expected_sheet]
                assert "DRAFT / INCOMPLETE" in str(workbook[expected_sheet]["A3"].value)
                values = [cell.value for cells in workbook[expected_sheet].iter_rows() for cell in cells]
                assert "Chưa có đơn giá / Not available" in values

            page.locator("#add-row").click()
            page.locator("#row-description").fill("Công việc khảo sát bổ sung")
            page.locator("#row-unit").fill("m")
            page.locator("#row-quantity").fill("2,5")
            page.locator("#row-save").click()
            page.wait_for_function(
                "() => (document.getElementById('estimate-rows').textContent || '').includes('Công việc khảo sát bổ sung')",
                timeout=20000,
            )

            created = page.request.post(
                "http://127.0.0.1:8000/api/admin/master-data/catalogue",
                data={
                    "code": "UX.TEST.001", "name_vi": "Công việc thử danh mục",
                    "name_en": "Catalogue UX test", "long_description": "Mô tả dài dùng làm bằng chứng đối chiếu",
                    "aliases": ["từ khóa thử", "bí danh"], "category": "Khảo sát", "unit": "m",
                    "unit_price": None, "status": "active", "actor": "uat",
                },
            )
            assert created.status == 201
            item_id = created.json()["item_id"]
            updated = page.request.patch(
                f"http://127.0.0.1:8000/api/admin/master-data/catalogue/{item_id}",
                data={
                    "name_vi": "Công việc thử danh mục đã sửa",
                    "long_description": "Mô tả chi tiết đã cập nhật để AI đối chiếu",
                    "aliases": ["từ khóa thử", "bí danh mới"],
                    "category": "Trắc địa", "unit": "m", "actor": "uat",
                },
            )
            assert updated.status == 200
            assert updated.json()["status"] == "updated"
            retired = page.request.delete(
                f"http://127.0.0.1:8000/api/admin/master-data/catalogue/{item_id}?actor=uat"
            )
            assert retired.status == 200
            assert retired.json()["physically_deleted"] is False

            page.locator("#catalogue-open").click()
            page.wait_for_function(
                "() => typeof catalogue !== 'undefined' && catalogue.some(x => x.code === 'UX.TEST.001')",
                timeout=20000,
            )
            page.locator("#catalogue-status").select_option("all")
            page.locator("#catalogue-search").fill("UX.TEST.001")
            assert page.locator("#catalogue-body", has_text="Ngừng sử dụng").count() == 1
            assert "đã sửa" in page.locator("#catalogue-body").text_content()
        finally:
            browser.close()


def test_review_fetch_error_shows_explicit_message(app_server):
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1440, "height": 1200})
        try:
            page.add_init_script("""
                const originalFetch = window.fetch.bind(window);
                window.fetch = (...args) => {
                    const url = String(args[0] || '');
                    if (url.includes('/api/review/')) return Promise.reject(new Error('review fetch failed'));
                    return originalFetch(...args);
                };
            """)
            page.goto("http://127.0.0.1:8000", wait_until="domcontentloaded")
            page.set_input_files("#pdf-input", [{
                "name": "sample.pdf", "mimeType": "application/pdf", "buffer": sample_pdf_bytes(),
            }])
            page.locator("#upload-btn").click()
            page.wait_for_function(
                "() => /lỗi|error/i.test(document.getElementById('status-bar').textContent || '')",
                timeout=30000,
            )
            assert "review fetch failed" in page.locator("#status-bar").text_content().lower()
        finally:
            browser.close()
