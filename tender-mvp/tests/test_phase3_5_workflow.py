import io
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "backend"))

import governance_store
from excel_exporter.exporter import generate_internal_workbook, generate_submission_workbook
from formula_config import approve_proposal, create_proposal, list_versions
from master_data.governance import add_project_override, overview


def _snapshot(*, complete: bool) -> dict:
    unresolved = 0 if complete else 1
    return {
        "snapshot_id": "calc-test-snapshot",
        "immutable": True,
        "job": {"id": "job-1", "filename": "Bang TIen Luong.pdf", "status": "needs_review"},
        "counts": {"source_rows": 95, "billable": 82, "non_billable": 13, "mapped": 81, "priced": 82 - unresolved,
                   "mandatory_unresolved": unresolved, "approval_evidence_covered": 82 - unresolved,
                   "approval_evidence_unresolved": unresolved},
        "formula_version": {
            "version_id": "calculation-v2.1", "effective_from": "2026-09-27",
            "tender": {"calculation": "validated tender"},
            "approval": {"calculation": "aggregate-first approval"},
        },
        "tender": {"method": "validated tender", "status": "COMPLETE" if complete else "INCOMPLETE",
                   "partial_subtotal": "1000", "official_total": "1000" if complete else None,
                   "section_subtotals": {"I": "1000"}},
        "approval": {"method": "aggregate-first", "status": "COMPLETE" if complete else "INCOMPLETE",
                     "partial_subtotal": "1100", "official_total": "1100" if complete else None,
                     "group_components": {"topography": {"FINAL": "1100"}}},
        "rows": [{
            "doc_order": 1, "row_id": "row-1", "row_number": "1", "row_type": "line_item",
            "section": "I", "description": "Verified work", "unit": "m", "quantity_raw": "2",
            "quantity": 2, "source_page": 1, "source_region": "[1,2]", "mapped_work_item": "CF.11620",
            "mapping_status": "mapped_and_priced", "tender_unit_price": "500", "tender_amount": "1000",
            "tender_formula_ref": "FORMULA-1", "price_source": "authoritative", "master_version": "v1",
            "evidence_conflicts": [] if complete else ["approval_resource_evidence_unavailable"],
            "review_history": [{"field": "quantity_raw", "old_value": "1", "new_value": "2", "user": "reviewer"}],
            "approval_detail": {"calculation_status": "resolved", "source_cells": ["Công trình!A16:AM16"],
                                "missing_dependencies": [], "resource_trace": [{"resource_type": "labour", "code": "NKS",
                                "description": "Labour", "quantity": "1", "coefficient": "1", "unit_price": "100",
                                "line_total": "100", "classification": "SOURCE_VERIFIED", "source_sheet": "Giá tháng",
                                "sheet_ref": "F1", "line_parity": "0"}]},
        }, {
            "doc_order": 2, "row_id": "II-II.3-II.3.2", "row_number": "II.3.2", "row_type": "metadata",
            "section": "II > II.3 > II.3.2", "description": "Cầu nhỏ, cầu trung, tỷ lệ 1/500: bình đồ, cắt dọc",
            "unit": "", "quantity_raw": "-", "quantity": None, "source_page": 2, "source_region": "[]",
            "source_origin": "source_document_inspection", "source_provenance": "Historical Azure fixture omitted it.",
            "azure_polygon_available": False, "mapped_work_item": None, "mapped_by": None,
            "mapping_status": "non_billable_metadata", "pricing_availability": "not_applicable",
            "price_display": "Not applicable", "tender_unit_price": None, "tender_amount": None,
            "evidence_conflicts": [], "review_history": [], "approval_detail": None,
        }],
        "unresolved_register": [] if complete else [{"row_id": "row-1", "row_number": "1", "description": "Verified work", "conflicts": ["approval_resource_evidence_unavailable"]}],
        "export_policy": {"internal_draft_allowed": True, "final_submission_allowed": complete,
                          "final_block_reasons": [] if complete else ["1 mandatory row unresolved"]},
    }


def test_canonical_reconciliation_is_95_source_rows_and_82_billable():
    evidence = json.loads((ROOT / "fixtures/canonical_source_contract_95.json").read_text())
    assert evidence["canonical_source_rows"] == 95
    assert evidence["historical_azure_rows"] == 94
    assert evidence["billable_rows"] == 82
    assert evidence["non_billable_rows"] == 13
    assert evidence["historical_pre_correction_billable_rows"] == 83
    assert evidence["nested_subsections_corrected"] == 1
    assert evidence["nested_subsection"]["stt"] == "II.3.1"
    assert evidence["recovered_structural_row"]["stt"] == "II.3.2"
    assert evidence["recovered_structural_row"]["azure_polygon_available"] is False


def test_internal_export_contains_audit_and_all_calculation_sheets():
    workbook = load_workbook(io.BytesIO(generate_internal_workbook(_snapshot(complete=False))), data_only=False)
    assert workbook.sheetnames == [
        "Summary", "Original BOQ & Corrections", "Detailed Resource Costs", "Tender Calculation",
        "Approval Calculation", "Formula & Sources", "Review History", "Unresolved Register",
    ]
    assert workbook["Summary"]["B3"].value == "calc-test-snapshot"
    assert workbook["Unresolved Register"].max_row == 4
    values = [cell.value for row in workbook["Original BOQ & Corrections"].iter_rows() for cell in row]
    assert "Not applicable" in values
    assert "source_document_inspection" in values


def test_final_export_is_blocked_when_mandatory_evidence_is_unresolved():
    with pytest.raises(ValueError, match="blocked"):
        generate_submission_workbook(_snapshot(complete=False))


def test_final_export_preserves_submission_contract_and_excludes_internal_data():
    workbook = load_workbook(io.BytesIO(generate_submission_workbook(_snapshot(complete=True))), data_only=False)
    assert workbook.sheetnames == ["Dự thầu"]
    sheet = workbook["Dự thầu"]
    assert sheet.page_setup.orientation == "portrait"
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.column_dimensions["F"].hidden is True
    assert str(sheet.print_area).startswith("'Dự thầu'!$A$1:$L$")
    assert sheet["J7"].value == "=ROUND(E7*I7,0)"
    values = " ".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row)
    assert "Approval" not in values
    assert "reviewer" not in values
    assert "audit" not in values.lower()


def test_master_candidates_remain_staging_and_conflicts_visible(monkeypatch, tmp_path):
    monkeypatch.setattr(governance_store, "RUNTIME_DIR", tmp_path)
    data = overview()
    assert data["candidate_count"] == 46
    assert len(data["resource_price_conflicts"]) == 3
    assert data["global_price_book_mutation_allowed"] is False
    assert all(candidate["scope"] == "staging_only" for candidate in data["candidates"])
    assert all(candidate["price_effective_date"] is None for candidate in data["candidates"])
    assert all(candidate["price_freshness"] == "effective_date_not_available" for candidate in data["candidates"])
    assert all(candidate["last_updated_by"] and candidate["last_updated_at"] for candidate in data["candidates"])
    assert data["effective_date_blockers"] == {
        "seeded_catalogue_missing": 30,
        "staged_candidates_missing": 46,
        "authoritative_dates_must_not_be_invented": True,
    }
    record = add_project_override({"project_id": "job-1", "resource_code": "A28.1000", "unit_price": "4480", "source_reference": "Công trình!A83", "effective_date": "2026-09-27"}, "tester")
    assert record["scope"] == "project_only"
    assert overview()["global_price_book_mutation_allowed"] is False

    undated = add_project_override({
        "project_id": "job-1", "resource_code": "UAT.NULL.DATE", "unit_price": "1",
        "source_reference": "controlled UAT evidence", "effective_date": None,
    }, "tester")
    assert undated["effective_date"] is None
    assert undated["price_freshness"] == "effective_date_not_available"


def test_formula_proposal_requires_preview_and_approval_before_use(monkeypatch, tmp_path):
    monkeypatch.setattr(governance_store, "RUNTIME_DIR", tmp_path)
    proposal = create_proposal({
        "name": "VAT preview", "effective_from": "2026-10-01", "category": "default",
        "basis": {"material": "100", "labour": "100", "machine": "100"},
        "changes": {"branch": "approval", "coefficients": {"vat_rate": "0.08"}},
    }, "proposer")
    assert proposal["status"] == "proposed"
    assert proposal["preview"]["before"] != proposal["preview"]["after"]
    assert len([item for item in list_versions() if item["status"] == "approved"]) == 1
    approved = approve_proposal(proposal["version_id"], "approver")
    assert approved["status"] == "approved"
    assert approved["audit"][-1]["event"] == "approved"
