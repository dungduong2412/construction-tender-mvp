import io
import json
import sys
from copy import deepcopy
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "backend"))

from excel_exporter.exporter import generate_internal_workbook, generate_submission_workbook
from pricing_availability import AUTHORIZED_ZERO, NOT_APPLICABLE, NOT_AVAILABLE, classify_price, display_price_state, price_freshness
from pricing_engine.aggregation import build_recursive_component_tree, build_recursive_money_tree


def test_null_and_authorized_zero_are_distinct():
    missing = classify_price(None, source_reference="approved schedule")
    unsupported_zero = classify_price("0", source_reference="approved schedule")
    authorized_zero = classify_price(
        "0", source_reference="signed decision 17", approval_status="approved",
        approved_by="commercial-director", zero_price_authorized=True,
    )
    assert missing.status == NOT_AVAILABLE and missing.value is None and not missing.available
    assert unsupported_zero.status == NOT_AVAILABLE and unsupported_zero.value is None
    assert authorized_zero.status == AUTHORIZED_ZERO and authorized_zero.value == Decimal("0")
    assert display_price_state(None, NOT_APPLICABLE) == "Not applicable"
    assert display_price_state(None, NOT_AVAILABLE) == "Not available"


def test_price_freshness_is_independent_of_availability():
    today = date(2026, 9, 28)
    assert price_freshness("2026-01-01", "2026-09-27", available=True, today=today) == "expired"
    assert price_freshness("2026-10-01", None, available=True, today=today) == "not_yet_effective"
    assert price_freshness("2026-01-01", None, available=True, today=today) == "current"
    assert price_freshness("2026-01-01", None, available=False, today=today) == "not_available"


def test_recursive_tender_aggregation_preserves_partial_and_zero_children():
    tree = build_recursive_money_tree([
        {"row_id": "1", "label": "priced", "path": "A > A.1", "amount": "10.4", "available": True},
        {"row_id": "2", "label": "authorized zero", "path": "A > A.1", "amount": "0", "available": True},
        {"row_id": "3", "label": "missing", "path": "A > A.2", "amount": None, "available": False},
    ], branch="tender", rounding_quantum="1")
    assert tree["status"] == "INCOMPLETE"
    assert tree["partial_subtotal"] == Decimal("10")
    assert tree["official_total"] is None
    assert tree["unresolved_descendants"] == 1
    section = tree["children"][0]
    assert section["official_total"] is None
    assert section["partial_subtotal"] == Decimal("10")


def test_recursive_approval_is_aggregate_first_at_every_parent():
    calls = []
    def calculate(group, components):
        calls.append((group, dict(components)))
        return components["material"] + components["labour"] + components["machine"]
    tree = build_recursive_component_tree([
        {"row_id": "1", "label": "one", "path": "A > A.1", "available": True,
         "components_by_group": {"survey": {"material": "0.4", "labour": "0.4", "machine": "0.4"}}},
        {"row_id": "2", "label": "two", "path": "A > A.1", "available": True,
         "components_by_group": {"survey": {"material": "0.4", "labour": "0.4", "machine": "0.4"}}},
    ], group_calculator=calculate, rounding_quantum="1")
    assert tree["official_total"] == Decimal("2")
    assert tree["unresolved_descendants"] == 0
    assert calls[-1][1]["material"] == Decimal("0.8")


def _incomplete_snapshot():
    return {
        "snapshot_id": "calc-business-rules", "immutable": True,
        "job": {"filename": "Bang TIen Luong.pdf", "status": "needs_review"},
        "counts": {"source_rows": 95, "billable": 82, "non_billable": 13, "mapped": 82, "priced": 66,
                   "mandatory_unresolved": 16, "approval_evidence_unresolved": 16},
        "formula_version": {"version_id": "calculation-v2.1", "tender": {}, "approval": {}},
        "tender": {"status": "INCOMPLETE", "partial_subtotal": "100", "official_total": None,
                   "method": "child aggregation", "hierarchy": {"depth": 0, "name": "Project",
                   "status": "INCOMPLETE", "partial_subtotal": "100", "official_total": None,
                   "unresolved_descendants": 16, "children": []}},
        "approval": {"status": "INCOMPLETE", "partial_subtotal": None, "official_total": None,
                     "method": "aggregate-first", "hierarchy": {"depth": 0, "name": "Project",
                     "status": "INCOMPLETE", "partial_subtotal": None, "official_total": None,
                     "unresolved_descendants": 16, "children": []}},
        "rows": [{"doc_order": 1, "row_number": "1", "row_type": "line_item", "section": "A",
                  "description": "Missing authoritative price", "unit": "m", "quantity_raw": "1",
                  "quantity": 1, "source_page": 1, "mapped_work_item": "KS.X",
                  "mapping_status": "mapped_price_unavailable", "pricing_availability": "not_available",
                  "tender_unit_price": None, "tender_amount": None, "evidence_conflicts": ["missing price"]}],
        "unresolved_register": [{"row_id": "1", "row_number": "1", "description": "Missing", "conflicts": ["missing price"]}],
        "export_policy": {"internal_draft_allowed": True, "final_submission_allowed": False,
                          "final_block_reasons": ["16 mandatory billable rows lack tender pricing"]},
    }


def test_internal_draft_says_not_available_and_final_stays_blocked():
    snapshot = _incomplete_snapshot()
    workbook = load_workbook(io.BytesIO(generate_internal_workbook(snapshot)), data_only=False)
    assert "DRAFT / INCOMPLETE" in workbook["Summary"]["A1"].value
    values = [cell.value for row in workbook["Original BOQ & Corrections"].iter_rows() for cell in row]
    assert "Not available" in values
    try:
        generate_submission_workbook(snapshot)
    except ValueError as exc:
        assert "blocked" in str(exc).lower()
    else:
        raise AssertionError("final export must be blocked")


def test_source_reconciliation_keeps_13_structural_separate_from_16_unpriced():
    contract = json.loads((ROOT / "fixtures/canonical_source_contract_95.json").read_text())
    assert contract["canonical_source_rows"] == 95
    assert contract["billable_rows"] == 82
    assert contract["historical_azure_rows"] == 94
    assert 95 - 82 == 13
    assert contract["unavailable_billable_rows"] == 16


def test_required_identity_and_freshness_surfaces_are_visible_in_ui():
    html = (ROOT / "frontend/index.html").read_text()
    for marker in ["price_effective_date", "price_last_updated_by", "mapping_last_updated_by",
                   "Pricing unavailable", "Not available", "Unresolved descendants"]:
        assert marker in html
