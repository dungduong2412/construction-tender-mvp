from __future__ import annotations

import copy
import json
import os
import uuid
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from integration_train1.mapper import MappingDecision, Train2SemanticMapper
from parser_adapter.provider_neutral import (
    HttpJsonProvider,
    ParserPayload,
    ParserProviderAuthError,
    ParserProviderContractError,
    ParserProviderError,
    ParserProviderTimeoutError,
    ProviderNeutralParserAdapter,
    RecordedFixtureProvider,
)
from pricing_engine.calculation_engine_v2 import CalculationEngineV2


FIXTURE_PATH = Path(__file__).resolve().parent.parent.parent / "fixtures" / "parser_api_recorded_train1.json"
REAL_RESPONSE_FIXTURE_PATH = Path(__file__).resolve().parent.parent.parent / "fixtures" / "parser_api_recorded_real_response.json"
UNIT_RULES_PATH = Path(__file__).resolve().parent.parent.parent / "master-data" / "unit_rules_v1.json"


class Train2PipelineError(Exception):
    pass


def _d(value: Any) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    return Decimal(str(value))


def _parse_quantity(raw: str) -> Decimal | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


@dataclass
class Train1Run:
    run_id: str
    source_document_id: str
    parsed_rows: list[dict[str, Any]]
    runtime: dict[str, Any]
    review: dict[str, Any]


class Train1Pipeline:
    def __init__(self):
        self.engine = CalculationEngineV2()
        self.mapper = Train2SemanticMapper(
            confidence_threshold=Decimal(str(os.getenv("TRAIN2_CONFIDENCE_THRESHOLD", "0.85")))
        )
        self.allow_zero_quantity = os.getenv("TRAIN2_ALLOW_ZERO_QUANTITY", "false").lower() in {"1", "true", "yes"}
        self._runs: dict[str, Train1Run] = {}
        self._unit_rules = self._load_unit_rules()

    def _load_unit_rules(self) -> list[dict[str, Any]]:
        if not UNIT_RULES_PATH.exists():
            return []
        try:
            payload = json.loads(UNIT_RULES_PATH.read_text(encoding="utf-8"))
            if not isinstance(payload, list):
                return []
            return payload
        except json.JSONDecodeError:
            return []

    async def start_from_fixture(self, fixture_path: Path | None = None) -> dict[str, Any]:
        provider = RecordedFixtureProvider(fixture_path or FIXTURE_PATH)
        adapter = ProviderNeutralParserAdapter(provider)
        parsed = await adapter.parse()
        return self.start_from_parsed_payload(parsed, live_integration_verified=False, integration_label="fixture")

    async def start_from_real_fixture(self) -> dict[str, Any]:
        provider = RecordedFixtureProvider(REAL_RESPONSE_FIXTURE_PATH)
        adapter = ProviderNeutralParserAdapter(provider)
        parsed = await adapter.parse()
        return self.start_from_parsed_payload(
            parsed,
            live_integration_verified=False,
            integration_label="recorded_real_response",
            integration_note="Live parser endpoint unavailable; using recorded real-response fixture.",
        )

    async def start_from_parser_api(
        self,
        pdf_bytes: bytes,
        endpoint: str,
        api_key: str | None = None,
        api_key_header: str = "Authorization",
        api_key_prefix: str = "Bearer",
        timeout_seconds: float = 30.0,
    ) -> dict[str, Any]:
        provider = HttpJsonProvider(
            endpoint=endpoint,
            api_key=api_key,
            api_key_header=api_key_header,
            api_key_prefix=api_key_prefix,
            timeout_seconds=timeout_seconds,
        )
        adapter = ProviderNeutralParserAdapter(provider)
        parsed = await adapter.parse(pdf_bytes)
        return self.start_from_parsed_payload(parsed, live_integration_verified=True, integration_label="live_api")

    def start_from_parsed_payload(
        self,
        parsed: dict[str, Any],
        *,
        live_integration_verified: bool = False,
        integration_label: str = "fixture",
        integration_note: str | None = None,
    ) -> dict[str, Any]:
        try:
            payload = ParserPayload.model_validate(parsed)
        except Exception as exc:
            raise Train2PipelineError(f"Invalid parser payload: {exc}") from exc

        runtime = self.engine.load_runtime_data()
        payload_dict = payload.model_dump()
        review = self._build_review(
            payload_dict,
            runtime,
            live_integration_verified=live_integration_verified,
            integration_label=integration_label,
            integration_note=integration_note,
        )
        run_id = str(uuid.uuid4())
        run = Train1Run(
            run_id=run_id,
            source_document_id=str(payload.document_id),
            parsed_rows=copy.deepcopy(payload_dict.get("rows", [])),
            runtime=runtime,
            review=review,
        )
        self._runs[run_id] = run
        return {"run_id": run_id, **review}

    def get_review(self, run_id: str) -> dict[str, Any]:
        run = self._runs.get(run_id)
        if run is None:
            raise KeyError(f"Run not found: {run_id}")
        return {"run_id": run_id, **run.review}

    def apply_manual_mapping(
        self,
        run_id: str,
        row_id: str,
        canonical_code: str,
        reason: str | None = None,
        unit_confirmed: bool = False,
        unit_correction: str | None = None,
    ) -> dict[str, Any]:
        run = self._runs.get(run_id)
        if run is None:
            raise KeyError(f"Run not found: {run_id}")

        rows = copy.deepcopy(run.parsed_rows)
        target = next((r for r in rows if str(r.get("row_id")) == row_id), None)
        if target is None:
            raise KeyError(f"Row not found: {row_id}")

        history = list(target.get("corrections") or [])
        history.append(
            {
                "action": "manual_mapping",
                "canonical_code": canonical_code,
                "reason": reason or "manual review override",
                "retained_evidence": list(target.get("evidence") or []),
                "source_unit_original": target.get("unit_raw"),
                "quantity_original": target.get("quantity_raw"),
                "unit_confirmed": unit_confirmed,
                "unit_correction": unit_correction,
            }
        )
        target["corrections"] = history
        target["manual_selected_code"] = canonical_code
        target["manual_override_evidence"] = reason or "manual_override"
        target["manual_unit_confirmed"] = bool(unit_confirmed)
        if unit_correction is not None:
            target["manual_unit_correction"] = str(unit_correction)

        run.parsed_rows = rows
        run.review = self._build_review(
            {"document_id": run.source_document_id, "rows": rows},
            run.runtime,
            live_integration_verified=bool(run.review.get("live_integration_verified")),
            integration_label=str(run.review.get("integration_label") or "fixture"),
            integration_note=run.review.get("integration_note"),
        )
        return {"run_id": run_id, **run.review}

    def mutate_runtime_prices(self, run_id: str, price_updates: dict[str, str]) -> dict[str, Any]:
        run = self._runs.get(run_id)
        if run is None:
            raise KeyError(f"Run not found: {run_id}")
        mutated_runtime = copy.deepcopy(run.runtime)
        price_book = mutated_runtime.get("resource_price_book") or {}
        for code, new_price in price_updates.items():
            if code not in price_book:
                continue
            price_book[code]["unit_price"] = str(_d(new_price))
        run.runtime = mutated_runtime
        run.review = self._build_review(
            {"document_id": run.source_document_id, "rows": run.parsed_rows},
            run.runtime,
            live_integration_verified=bool(run.review.get("live_integration_verified")),
            integration_label=str(run.review.get("integration_label") or "fixture"),
            integration_note=run.review.get("integration_note"),
        )
        return {"run_id": run_id, **run.review}

    def mutate_approval_group_rules(self, run_id: str, group_key: str, updates: dict[str, str]) -> dict[str, Any]:
        run = self._runs.get(run_id)
        if run is None:
            raise KeyError(f"Run not found: {run_id}")
        mutated_runtime = copy.deepcopy(run.runtime)
        groups = mutated_runtime.setdefault("approval_rule_groups", {})
        group = groups.setdefault(group_key, {})
        rules = group.setdefault("approval_rules", {})
        rules.update(updates)
        run.runtime = mutated_runtime
        run.review = self._build_review(
            {"document_id": run.source_document_id, "rows": run.parsed_rows},
            run.runtime,
            live_integration_verified=bool(run.review.get("live_integration_verified")),
            integration_label=str(run.review.get("integration_label") or "fixture"),
            integration_note=run.review.get("integration_note"),
        )
        return {"run_id": run_id, **run.review}

    def export_excel(self, run_id: str) -> bytes:
        run = self._runs.get(run_id)
        if run is None:
            raise KeyError(f"Run not found: {run_id}")
        review = run.review

        wb = Workbook()
        ws = wb.active
        ws.title = "Train1-Review"

        ws.append([
            "row_id",
            "description_vi",
            "source_unit",
            "master_unit",
            "source_quantity",
            "converted_quantity",
            "quantity_factor",
            "mapping_status",
            "canonical_code",
            "loaded_unit_price",
            "tender_unit_price",
            "tender_extension",
            "confidence",
            "evidence",
        ])
        for row in review["rows"]:
            ws.append(
                [
                    row["row_id"],
                    row["description_vi"],
                    row.get("source_unit") or row["unit_raw"],
                    row.get("master_unit"),
                    row.get("source_quantity"),
                    float(_d(row["converted_quantity"])) if row.get("converted_quantity") not in {None, ""} else None,
                    float(_d(row["quantity_factor"])) if row.get("quantity_factor") not in {None, ""} else None,
                    row["mapping_status"],
                    row.get("canonical_code"),
                    float(_d(row["loaded_unit_price"])) if row.get("loaded_unit_price") is not None else None,
                    float(_d(row["tender_unit_price"])) if row.get("tender_unit_price") is not None else None,
                    float(_d(row["tender_extension"])) if row.get("tender_extension") is not None else None,
                    row.get("confidence"),
                    row.get("mapping_evidence"),
                ]
            )

        ws2 = wb.create_sheet("Train1-Summary")
        ws2.append(["status", review["status"]])
        ws2.append(["tender_partial_subtotal", float(_d(review["tender_partial_subtotal"]))])
        ws2.append(["approval_partial_subtotal", float(_d(review["approval_partial_subtotal"]))])
        ws2.append(["official_tender_total", float(_d(review["official_tender_total"])) if review["official_tender_total"] is not None else None])
        ws2.append(["official_approval_total", float(_d(review["official_approval_total"])) if review["official_approval_total"] is not None else None])

        ws3 = wb.create_sheet("Train1-Unresolved")
        ws3.append(["row_id", "mapping_status", "reason", "missing_dependencies"])
        for row in review["rows"]:
            if row["mapping_status"] != "resolved":
                ws3.append(
                    [
                        row["row_id"],
                        row["mapping_status"],
                        row.get("reason"),
                        ", ".join(row.get("missing_dependencies") or []),
                    ]
                )

        import io

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def _build_review(
        self,
        parsed: dict[str, Any],
        runtime: dict[str, Any],
        *,
        live_integration_verified: bool,
        integration_label: str,
        integration_note: str | None,
    ) -> dict[str, Any]:
        work_master = runtime.get("work_item_master") or {}

        rows_out: list[dict[str, Any]] = []
        tender_partial = Decimal("0")
        unresolved_required = 0
        required_rows = 0
        blocked_rows: list[dict[str, Any]] = []

        approval_group_accumulator: dict[str, dict[str, Any]] = {}

        for index, row in enumerate(parsed.get("rows", [])):
            row_id = str(row.get("row_id") or f"row-{index+1}")
            row_type = str(row.get("row_type") or "line_item")
            description_vi = str(row.get("description_vi") or "")
            unit_raw = str(row.get("unit_raw") or "")
            quantity_raw = str(row.get("quantity_raw") or "")
            quantity = _parse_quantity(quantity_raw)
            evidence = list(row.get("evidence") or [])
            suggestions = list(row.get("ai_suggestions") or [])
            corrections = list(row.get("corrections") or [])
            manual_selected_code = str(row.get("manual_selected_code") or "").strip()
            manual_unit_confirmed = bool(row.get("manual_unit_confirmed"))
            manual_unit_correction = row.get("manual_unit_correction")

            if row_type == "line_item":
                required_rows += 1

            decision: MappingDecision = self.mapper.decide(
                row_id=row_id,
                row_type=row_type,
                description_vi=description_vi,
                unit_raw=unit_raw,
                quantity_raw=quantity_raw,
                quantity=quantity,
                suggestions=suggestions,
                work_master=work_master,
                unit_rules=self._unit_rules,
                manual_selected_code=manual_selected_code,
                manual_unit_confirmed=manual_unit_confirmed,
                manual_unit_correction=str(manual_unit_correction) if manual_unit_correction is not None else None,
                allow_zero_quantity=self.allow_zero_quantity,
            )

            mapping_status = decision.mapping_status
            canonical_code = decision.canonical_code
            confidence = decision.confidence
            reason = decision.reason
            missing_dependencies = list(decision.missing_dependencies)
            loaded_unit_price = None
            tender_unit_price = None
            tender_extension = None
            approval_unit_price = None
            converted_quantity = None
            quantity_factor = decision.quantity_factor
            source_unit = decision.source_unit
            master_unit = decision.master_unit

            if mapping_status == "resolved" and canonical_code is not None and quantity is not None:
                item_result = self.engine.calculate_work_item_from_runtime(canonical_code, runtime)
                if item_result.get("calculation_status") != "resolved":
                    mapping_status = "unsupported_blocking_item"
                    reason = "Mapped item has unresolved required dependencies"
                    missing_dependencies = list(item_result.get("missing_dependencies") or [])
                else:
                    work_item = work_master.get(canonical_code) or {}
                    loaded_unit = item_result["loaded_unit_price"]
                    tender_unit = item_result["tender_rounded_unit_price"]
                    loaded_unit_price = str(loaded_unit)
                    tender_unit_price = str(tender_unit)
                    effective_quantity = (quantity * quantity_factor).quantize(Decimal("0.0001"))
                    converted_quantity = str(effective_quantity)
                    tender_extension_value = (tender_unit * effective_quantity).quantize(Decimal("1"))
                    tender_extension = str(tender_extension_value)
                    tender_partial += tender_extension_value

                    # Aggregate-first approval accumulator (do not sum per-row loaded approval finals).
                    category = str(work_item.get("category") or "default")
                    approval_group = str(work_item.get("approval_rule_group") or "") or self.engine._approval_group_for_item(work_item)
                    bucket = approval_group_accumulator.setdefault(
                        approval_group,
                        {
                            "category": category,
                            "material": Decimal("0"),
                            "labour": Decimal("0"),
                            "machine": Decimal("0"),
                        },
                    )
                    bucket["material"] += item_result["direct_material_total"] * effective_quantity
                    bucket["labour"] += item_result["direct_labour_total"] * effective_quantity
                    bucket["machine"] += item_result["direct_machine_total"] * effective_quantity

                    # Per-row display only.
                    approval_rules = self.engine.approval_rules_for(category, work_item, runtime=runtime, approval_group=approval_group)
                    approval_components = self.engine._compute_loaded_components(
                        item_result["direct_material_total"],
                        item_result["direct_labour_total"],
                        item_result["direct_machine_total"],
                        approval_rules,
                    )
                    approval_unit_price = str(approval_components["FINAL"])

            if row_type == "line_item" and mapping_status != "resolved":
                unresolved_required += 1
                blocked_rows.append({"row_id": row_id, "mapping_status": mapping_status, "reason": reason})

            rows_out.append(
                {
                    "row_id": row_id,
                    "row_type": row_type,
                    "description_vi": description_vi,
                    "unit_raw": unit_raw,
                    "quantity_raw": quantity_raw,
                    "quantity": str(quantity) if quantity is not None else None,
                    "source_quantity": quantity_raw,
                    "converted_quantity": converted_quantity,
                    "quantity_factor": str(quantity_factor),
                    "source_unit": source_unit,
                    "master_unit": master_unit,
                    "evidence": evidence,
                    "corrections": corrections,
                    "mapping_status": mapping_status,
                    "canonical_code": canonical_code,
                    "confidence": confidence,
                    "mapping_evidence": decision.mapping_evidence,
                    "candidates": decision.candidates,
                    "reason": reason,
                    "missing_dependencies": missing_dependencies,
                    "loaded_unit_price": loaded_unit_price,
                    "tender_unit_price": tender_unit_price,
                    "tender_extension": tender_extension,
                    "approval_unit_price": approval_unit_price,
                }
            )

        approval_partial = Decimal("0")
        approval_group_components: dict[str, dict[str, str]] = {}
        for group_key, totals in approval_group_accumulator.items():
            category = str(totals["category"])
            rules = self.engine.approval_rules_for(category, runtime=runtime, approval_group=group_key)
            components = self.engine._compute_loaded_components(
                totals["material"],
                totals["labour"],
                totals["machine"],
                rules,
            )
            approval_partial += components["FINAL"]
            approval_group_components[group_key] = {
                "category": category,
                "material": str(totals["material"]),
                "labour": str(totals["labour"]),
                "machine": str(totals["machine"]),
                "final": str(components["FINAL"]),
            }

        complete = required_rows > 0 and unresolved_required == 0
        return {
            "source_document_id": str(parsed.get("document_id") or "fixture-document"),
            "required_row_count": required_rows,
            "resolved_required_row_count": required_rows - unresolved_required,
            "status": "COMPLETE" if complete else "INCOMPLETE",
            "official_tender_total": str(tender_partial) if complete else None,
            "official_approval_total": str(approval_partial) if complete else None,
            "tender_partial_subtotal": str(tender_partial),
            "approval_partial_subtotal": str(approval_partial),
            "approval_group_components": approval_group_components,
            "blocked_rows": blocked_rows,
            "live_integration_verified": live_integration_verified,
            "integration_label": integration_label,
            "integration_note": integration_note,
            "rows": rows_out,
        }


PIPELINE = Train1Pipeline()


__all__ = [
    "PIPELINE",
    "Train1Pipeline",
    "Train2PipelineError",
    "ParserProviderError",
    "ParserProviderAuthError",
    "ParserProviderTimeoutError",
    "ParserProviderContractError",
]
